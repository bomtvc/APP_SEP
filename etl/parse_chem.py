# -*- coding: utf-8 -*-
"""Đọc "list hóa chất total 2026.xlsx" -> data/chem_master.csv.

Đây là DANH MỤC THAM CHIẾU DUY NHẤT của hệ thống: mọi mã không có ở đây đều bị
loại khỏi thống kê (xem config.RESTRICT_TO_MASTER_LIST).

Vì vậy script dừng ngay nếu phát hiện mã trùng lặp — danh mục tham chiếu mà có
hai dòng cho cùng một CODE thì mọi phép join phía sau đều nhân đôi số liệu.
"""
import csv
import os
import re
import sys

sys.stdout = __import__('io').TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl
from openpyxl.utils import column_index_from_string as ci

import config

CAS_RE = re.compile(r'(\d{2,7}-\d{2}-\d)')

CORRECTIONS_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'cas_corrections.csv')

# Bảng phân loại GHS: (tên cột đầu ra, cột "CẤP ĐỘ", cột "HS CODE" tức mã H).
# Tám nhóm đầu chiếm hai cột liền nhau. Nhóm "NGUY HẠI/NGUY HIỂM KHÁC" chỉ có MỘT
# cột (W) và không có ô mã H — cột X kế bên là đường dẫn MSDS, không phải mã H.
HAZARD_COLUMNS = [
    ('explosive', 'G', 'H'), ('flammable', 'I', 'J'), ('oxidizing', 'K', 'L'),
    ('gas_pressure', 'M', 'N'), ('corrosive', 'O', 'P'), ('toxic', 'Q', 'R'),
    ('health', 'S', 'T'), ('env', 'U', 'V'), ('other', 'W', None),
]


def load_corrections():
    """Bảng đính chính số CAS của danh mục, đối chiếu với MSDS.

    Không sửa thẳng file Excel của bộ phận vận hành: đính chính để ở đây nên vừa
    tái lập được sau mỗi lần chạy, vừa có lý do và nguồn kèm theo. Khi file nguồn
    đã được sửa, dòng đính chính tự báo là không còn cần thiết và có thể xoá.
    """
    if not os.path.exists(CORRECTIONS_FILE):
        return []
    with open(CORRECTIONS_FILE, encoding='utf-8-sig') as fh:
        return [r for r in csv.DictReader(fh) if r.get('code')]


def apply_corrections(rows, corrections):
    by_code = {r['code']: r for r in rows}
    applied, stale, unmatched = [], [], []
    for c in corrections:
        rec = by_code.get(c['code'])
        if rec is None:
            unmatched.append((c['code'], c['cas'], 'khong co ma nay trong danh muc'))
            continue
        cas_list = [x for x in rec['cas'].split(';') if x]
        if c['hanh_dong'] == 'bo':
            if c['cas'] not in cas_list:
                stale.append((c['code'], c['cas']))
                continue
            rec['cas_raw_goc'] = rec['cas_raw']
            rec['cas'] = ';'.join(x for x in cas_list if x != c['cas'])
            rec['cas_raw'] = '; '.join(
                p.strip() for p in rec['cas_raw'].split(';')
                if c['cas'] not in p)
            applied.append((c['code'], c['cas'], 'bo'))
        elif c['hanh_dong'] == 'them':
            if c['cas'] in cas_list:
                stale.append((c['code'], c['cas']))
                continue
            rec['cas_raw_goc'] = rec['cas_raw']
            rec['cas'] = ';'.join(cas_list + [c['cas']])
            rec['cas_raw'] = rec['cas_raw'] + '; ' + c['cas']
            applied.append((c['code'], c['cas'], 'them'))
        else:
            unmatched.append((c['code'], c['cas'], 'hanh_dong khong hop le: ' + c['hanh_dong']))
    return applied, stale, unmatched


def main(out_path):
    ws = openpyxl.load_workbook(config.SRC_CHEM_LIST, data_only=True)[config.SRC_CHEM_LIST_SHEET]

    rows = []
    for r in range(4, ws.max_row + 1):
        code = ws.cell(r, 2).value
        if not code:
            continue
        code = str(code).strip()
        cas_raw = str(ws.cell(r, 4).value or '')
        cas = list(dict.fromkeys(CAS_RE.findall(cas_raw)))

        rec = {
            'stt': ws.cell(r, 1).value,
            'code': code,
            'name': str(ws.cell(r, 3).value or '').strip(),
            'cas_raw': cas_raw.replace('\n', '; '),
            'cas': ';'.join(cas),
            'max_stock_kg': ws.cell(r, 5).value,
            'state': ws.cell(r, 6).value,
            'msds_path': ws.cell(r, 24).value,
            'src_row': r,
            'cas_raw_goc': '',
        }
        for name, cat_col, h_col in HAZARD_COLUMNS:
            rec[name + '_cat'] = ws.cell(r, ci(cat_col)).value
            rec[name + '_h'] = ws.cell(r, ci(h_col)).value if h_col else None
        rows.append(rec)

    # --- kiểm tra toàn vẹn: danh mục tham chiếu không được có mã trùng ---
    seen = {}
    dups = []
    for rec in rows:
        if rec['code'] in seen:
            dups.append((rec['code'], seen[rec['code']], rec['src_row']))
        else:
            seen[rec['code']] = rec['src_row']
    if dups:
        for code, r1, r2 in dups:
            print('  TRÙNG LẶP: %s ở dòng %d và %d' % (code, r1, r2))
        sys.exit('DỪNG: danh mục tham chiếu có %d mã trùng lặp. Sửa file nguồn rồi chạy lại.' % len(dups))

    applied, stale, unmatched = apply_corrections(rows, load_corrections())

    with open(out_path, 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)

    # --- đối chiếu MSDS theo CODE ---
    files = {os.path.splitext(f)[0].strip().upper() for f in os.listdir(config.SRC_MSDS_DIR)}
    missing = [r['code'] for r in rows if r['code'].upper() not in files]

    print('Danh mục tham chiếu: %d mã, không có mã trùng lặp.' % len(rows))
    for code, cas, act in applied:
        print('  đính chính CAS: %s %s %s (theo MSDS)' % (code, 'bỏ' if act == 'bo' else 'thêm', cas))
    for code, cas in stale:
        print('  ! đính chính %s / %s không còn cần — file nguồn đã đúng, có thể xoá dòng này'
              % (code, cas))
    for code, cas, why in unmatched:
        print('  ! đính chính %s / %s bị bỏ qua: %s' % (code, cas, why))
    print('Không có mã CAS: %d mã' % sum(1 for r in rows if not r['cas']))
    print('MSDS khớp theo CODE: %d/%d' % (len(rows) - len(missing), len(rows)))
    if missing:
        print('  Thiếu MSDS:', ', '.join(missing))


if __name__ == '__main__':
    main(sys.argv[1])
